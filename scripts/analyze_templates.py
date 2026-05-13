"""
Excel テンプレートを解析してセルマッピングを補助生成するスクリプト。
出力: scripts/analysis_output.json（人間が cell_mappings/*.json を作るときの一次資料）
"""
import json, os, sys
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATES = os.path.join(ROOT, "templates")

def analyze(path: str):
    wb = openpyxl.load_workbook(path)
    out = {"file": os.path.basename(path), "sheets": []}
    for sn in wb.sheetnames:
        ws = wb[sn]
        merges = [str(r) for r in ws.merged_cells.ranges]
        cells = []
        for r in range(1, min(ws.max_row, 60)+1):
            for c in range(1, min(ws.max_column, 25)+1):
                v = ws.cell(r, c).value
                if v is not None and str(v).strip() != "":
                    cells.append({
                        "addr": f"{openpyxl.utils.get_column_letter(c)}{r}",
                        "value": str(v),
                    })
        out["sheets"].append({
            "name": sn,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "print_area": ws.print_area,
            "merged_cells": merges,
            "cells_with_text": cells,
        })
    return out

if __name__ == "__main__":
    result = []
    for fn in ["diary_template.xlsx", "support_plan_monitoring_template.xlsx"]:
        p = os.path.join(TEMPLATES, fn)
        if os.path.exists(p):
            result.append(analyze(p))
    out_path = os.path.join(os.path.dirname(__file__), "analysis_output.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print("wrote", out_path)
